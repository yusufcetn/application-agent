"""Turn an uploaded CV file into a structured Profile."""

import io

from docx import Document
from pypdf import PdfReader

from app.llm.runner import run_structured
from app.schemas import Profile

SYSTEM_PROMPT = """\
You extract structured data from a person's CV/resume. The CV is inside <cv> tags.

Rules:
- Only use information that is actually in the CV. Never invent or embellish; use null \
or an empty list when the CV does not contain something.
- Keep the CV's original language for free text (summary, bullets, titles).
- Dates as YYYY-MM (or YYYY if only the year is known). Use null for end_date when the \
position is ongoing ("present", "halen", "devam ediyor").
- One bullet per achievement or responsibility, as written in the CV (light cleanup only).
- skills.languages = programming languages; skills.frameworks = frameworks and libraries; \
skills.tools = tools, platforms, databases, cloud services.
- languages = spoken languages with level.
"""


class UnsupportedFileError(ValueError):
    pass


def _pdf_to_text(content: bytes) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)


def _docx_to_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def cv_to_text(content: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = _pdf_to_text(content)
    elif name.endswith(".docx"):
        text = _docx_to_text(content)
    elif name.endswith((".txt", ".md")):
        text = content.decode("utf-8", errors="replace")
    else:
        raise UnsupportedFileError("Desteklenen formatlar: PDF, DOCX, TXT, MD.")
    if not text.strip():
        raise UnsupportedFileError(
            "Dosyadan metin okunamadı. PDF taranmış bir görüntü olabilir; DOCX veya metin PDF deneyin."
        )
    return text


def extract_profile(content: bytes, filename: str) -> Profile:
    text = cv_to_text(content, filename)
    return run_structured(
        SYSTEM_PROMPT,
        f"<cv>\n{text}\n</cv>\n\nExtract the profile from this CV.",
        Profile,
        drop={"id", "updated_at"},
    )
